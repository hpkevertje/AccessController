"""Webapp: leden, logboek, configuratie, status en manuele bediening."""
import io
import logging
import os
from contextlib import asynccontextmanager
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

import openpyxl
from fastapi import Body, FastAPI, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from starlette.concurrency import run_in_threadpool

from . import auth, backup, db, devices, members, settings, sharepoint
from .engine import Engine

logging.basicConfig(level=os.environ.get("LOG_LEVEL", "INFO"),
                    format="%(asctime)s %(levelname)s %(name)s: %(message)s")
log = logging.getLogger("accesscontrol")

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LOCAL_TZ = ZoneInfo(os.environ.get("TZ", "Europe/Brussels"))
STATIC_DIR = os.path.join(os.path.dirname(__file__), "static")
engine: Engine | None = None


@asynccontextmanager
async def lifespan(app: FastAPI):
    global engine
    db.init()
    db.log_event("systeem", result="info", reason="Server gestart")
    if not auth.configured():
        log.error("Cloudflare Access niet geconfigureerd: alle verzoeken worden geweigerd.")
    engine = Engine()
    if os.environ.get("ENGINE_DISABLED", "false").lower() not in ("1", "true", "yes"):
        engine.start()
    yield
    engine.stop()


app = FastAPI(title="AccessController", lifespan=lifespan, docs_url=None, redoc_url=None)


@app.middleware("http")
async def require_cloudflare(request: Request, call_next):
    if request.url.path == "/healthz":
        return await call_next(request)
    try:
        request.state.user = await run_in_threadpool(auth.current_user, request)
    except HTTPException as e:
        if request.url.path.startswith("/api/"):
            return JSONResponse({"detail": e.detail}, status_code=e.status_code)
        return Response(f"<h1>Geen toegang</h1><p>{e.detail}</p>", status_code=e.status_code,
                        media_type="text/html")
    return await call_next(request)


def user(request: Request) -> str:
    return request.state.user


def region() -> str:
    return engine.cfg["phone_region"]


def _validation(e: members.ValidationError):
    return JSONResponse({"detail": "Validatie mislukt", "errors": e.errors}, status_code=422)


def _stamp() -> str:
    return datetime.now(LOCAL_TZ).strftime("%Y%m%d-%H%M")


def _local(ts: str | None) -> str:
    if not ts:
        return ""
    dt = datetime.strptime(ts, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=timezone.utc)
    return dt.astimezone(LOCAL_TZ).strftime("%Y-%m-%d %H:%M:%S")


@app.get("/healthz")
def healthz():
    return {"ok": True}


@app.get("/api/me")
def me(request: Request):
    return {"email": user(request), "auth_disabled": auth.AUTH_DISABLED}


def _sms_unread() -> dict:
    seen = db.get_setting("sms_seen_id", 0)
    row = db.query_one("SELECT COUNT(*) AS n, MAX(id) AS last FROM events WHERE category='sms' "
                       "AND IFNULL(reason, '') != 'Lege SMS ontvangen' AND id > ?", (seen,))
    newest = None
    if row["n"]:
        newest = db.query_one("SELECT ts, device, actor, member_name, detail FROM events WHERE id=?",
                              (row["last"],))
    return {"count": row["n"], "newest": newest}


@app.get("/api/status")
def status():
    return {**engine.status(), "sms_unread": _sms_unread()}


@app.post("/api/sms/seen")
def sms_seen(request: Request):
    """Markeert alle SMS als gelezen, voor alle admins."""
    unread = _sms_unread()
    last = db.query_one("SELECT MAX(id) AS id FROM events WHERE category='sms'")["id"] or 0
    db.set_setting("sms_seen_id", last)
    if unread["count"]:
        db.log_event("admin", actor=user(request), result="ok", reason="SMS gelezen",
                     detail=f"{unread['count']} nieuwe SMS als gelezen gemarkeerd")
    return {"ok": True}


# ---------- leden ----------

@app.get("/api/members")
def list_members():
    return members.list_members()


@app.post("/api/members")
def create_member(request: Request, data: dict = Body(...)):
    try:
        return members.create(data, region(), user(request))
    except members.ValidationError as e:
        return _validation(e)


@app.put("/api/members/{member_id}")
def update_member(member_id: int, request: Request, data: dict = Body(...)):
    try:
        return members.update(member_id, data, region(), user(request))
    except members.ValidationError as e:
        return _validation(e)
    except KeyError:
        raise HTTPException(404, "Lid niet gevonden")


@app.delete("/api/members/{member_id}")
def delete_member(member_id: int, request: Request):
    try:
        members.delete(member_id, user(request))
    except KeyError:
        raise HTTPException(404, "Lid niet gevonden")
    return {"ok": True}


@app.get("/api/members/export.xlsx")
def export_xlsx(request: Request):
    db.log_event("admin", actor=user(request), result="ok", reason="Export leden")
    return Response(members.export_xlsx(), media_type=XLSX,
                    headers={"Content-Disposition": f'attachment; filename="leden-{_stamp()}.xlsx"'})


@app.post("/api/members/import")
async def import_members(request: Request, file: UploadFile = File(...)):
    content = await file.read()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, "Bestand is groter dan 5 MB.")
    try:
        return members.import_file(file.filename or "", content, region(), user(request))
    except members.ValidationError as e:
        db.log_event("admin", actor=user(request), result="fout", reason="Import leden geweigerd",
                     detail=f"{file.filename}: {len(e.errors)} fout(en). " + " | ".join(e.errors[:20]))
        return _validation(e)


# ---------- logboek ----------

def _event_filter(frm, to, category, device, result, q):
    where, params = [], []
    if frm:
        where.append("ts >= ?")
        params.append(frm)
    if to:
        where.append("ts <= ?")
        params.append(to)
    if device:
        where.append("device_id = ?")
        params.append(int(device))
    for col, val in (("category", category), ("result", result)):
        if val:
            where.append(f"{col} = ?")
            params.append(val)
    if q:
        like = f"%{q.strip()}%"
        where.append("(member_name LIKE ? OR member_no LIKE ? OR phone LIKE ? OR actor LIKE ? "
                     "OR reason LIKE ? OR detail LIKE ?)")
        params += [like] * 6
    return (" WHERE " + " AND ".join(where)) if where else "", params


@app.get("/api/events")
def list_events(frm: str | None = Query(None, alias="from"), to: str | None = None,
                category: str | None = None, device: str | None = None, result: str | None = None,
                q: str | None = None, page: int = 1, page_size: int = 100):
    page_size = max(1, min(page_size, 500))
    where, params = _event_filter(frm, to, category, device, result, q)
    total = db.query_one(f"SELECT COUNT(*) AS n FROM events{where}", params)["n"]
    rows = db.query(f"SELECT * FROM events{where} ORDER BY ts DESC, id DESC LIMIT ? OFFSET ?",
                    params + [page_size, (max(page, 1) - 1) * page_size])
    return {"total": total, "page": page, "page_size": page_size, "items": rows}


@app.get("/api/events/export.xlsx")
def export_events(request: Request, frm: str | None = Query(None, alias="from"), to: str | None = None,
                  category: str | None = None, device: str | None = None, result: str | None = None,
                  q: str | None = None):
    where, params = _event_filter(frm, to, category, device, result, q)
    rows = db.query(f"SELECT * FROM events{where} ORDER BY ts DESC, id DESC", params)
    cols = [("ts", "Tijdstip", 20), ("category", "Soort", 10), ("device", "Toestel", 9), ("actor", "Door", 26),
            ("member_no", "Lidnr", 9), ("member_name", "Naam", 24), ("phone", "Telefoon", 16),
            ("result", "Resultaat", 11), ("reason", "Omschrijving", 32), ("detail", "Detail", 60),
            ("confirmed_at", "Bevestigd door relais", 20)]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Logboek"
    ws.append([h for _, h, _ in cols])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([_local(r[c]) if c in ("ts", "confirmed_at") else (r[c] or "")
                   for c, _, _ in cols])
    for i, (_, _, w) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    db.log_event("admin", actor=user(request), result="ok", reason="Export logboek (Excel)",
                 detail=f"{len(rows)} regels")
    return Response(buf.getvalue(), media_type=XLSX,
                    headers={"Content-Disposition": f'attachment; filename="logboek-{_stamp()}.xlsx"'})


# ---------- configuratie ----------

@app.get("/api/config")
def get_config():
    return settings.public(engine.cfg)


@app.put("/api/config")
def put_config(request: Request, data: dict = Body(...)):
    old = engine.cfg
    new = dict(data)
    mqtt_in = dict(new.get("mqtt") or {})
    mqtt_in.pop("password_set", None)
    if not mqtt_in.get("password") and not mqtt_in.pop("clear_password", False):
        mqtt_in["password"] = old["mqtt"]["password"]       # leeg = ongewijzigd
    mqtt_in.pop("clear_password", None)
    new["mqtt"] = mqtt_in
    try:
        cfg = settings.save(new)
    except settings.ConfigError as e:
        return JSONResponse({"detail": "Validatie mislukt", "errors": e.args[0]}, status_code=422)
    changes = settings.diff(old, cfg)
    if changes:
        db.log_event("admin", actor=user(request), result="ok", reason="Configuratie gewijzigd",
                     detail="; ".join(changes))
        engine.reload(cfg)
    return settings.public(cfg)


# ---------- toestellen ----------

def _device_error(e: devices.DeviceError):
    return JSONResponse({"detail": "Validatie mislukt", "errors": e.errors}, status_code=422)


@app.get("/api/devices")
def list_devices():
    return devices.list_devices()


@app.post("/api/devices")
def create_device(request: Request, data: dict = Body(...)):
    grant = data.get("grant_all")
    try:
        dev = devices.create(data, None if grant is None else bool(grant), user(request))
    except devices.DeviceError as e:
        return _device_error(e)
    engine.reload_devices()
    return dev


@app.put("/api/devices/{device_id}")
def update_device(device_id: int, request: Request, data: dict = Body(...)):
    try:
        dev = devices.update(device_id, data, user(request))
    except devices.DeviceError as e:
        return _device_error(e)
    except KeyError:
        raise HTTPException(404, "Toestel niet gevonden")
    engine.reload_devices()
    return dev


@app.delete("/api/devices/{device_id}")
def delete_device(device_id: int, request: Request):
    try:
        devices.delete(device_id, user(request))
    except KeyError:
        raise HTTPException(404, "Toestel niet gevonden")
    engine.reload_devices()
    return {"ok": True}


@app.post("/api/devices/{device_id}/{action}")
def device_action(device_id: int, action: str, request: Request):
    if action not in ("on", "off"):
        raise HTTPException(404, "Onbekende actie")
    try:
        cmd_id = engine.manual(device_id, action, user(request))
    except KeyError:
        raise HTTPException(404, "Toestel niet gevonden")
    except ValueError as e:
        raise HTTPException(409, str(e))
    except RuntimeError as e:
        raise HTTPException(503, f"Commando niet verzonden: {e}")
    return {"ok": True, "cmd_id": cmd_id}


# ---------- back-up ----------

@app.get("/api/backup/status")
def backup_status():
    return {**backup.status(), "sharepoint_secret_set": bool(sharepoint.secret())}


@app.post("/api/backup/run")
def backup_run(request: Request):
    return backup.run(engine.cfg, actor=user(request))


@app.get("/api/backup")
def backup_download(request: Request):
    name, data = backup.make()
    db.log_event("admin", actor=user(request), result="ok", reason="Back-up gedownload", detail=name)
    return Response(data, media_type="application/gzip",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.post("/api/backup/restore")
async def backup_restore(request: Request, file: UploadFile = File(...)):
    content = await file.read()
    if len(content) > 200 * 1024 * 1024:
        raise HTTPException(413, "Bestand is groter dan 200 MB.")
    try:
        result = await run_in_threadpool(backup.restore, content, engine.cfg, user(request), file.filename or "")
    except backup.BackupError as e:
        db.log_event("admin", actor=user(request), result="fout", reason="Terugzetten back-up geweigerd",
                     detail=f"{file.filename}: {e}")
        return JSONResponse({"detail": str(e), "errors": [str(e)]}, status_code=422)
    # Configuratie en toestellen komen nu uit de teruggezette database.
    engine.reload(settings.load())
    engine.reload_devices()
    return result


app.mount("/", StaticFiles(directory=STATIC_DIR, html=True), name="static")
