"""Leden: validatie, CRUD en import/export (Excel, zelfde formaat in en uit).

Toegang per toestel staat in member_access. In de API: "access": {"<toestel-id>": true/false}.
"""
import io
import re

import openpyxl
from openpyxl.styles import Font

from . import db, devices
from .phone import PhoneError, normalize

FIELDS = ("member_no", "name", "phone", "note")
LABELS = {"member_no": "Lidnr", "name": "Naam", "phone": "Telefoon", "note": "Opmerking"}

# Herkende kolomkoppen bij import (kleine letters, zonder spaties/leestekens).
HEADER_ALIASES = {
    "member_no": {"lidnr", "lidnummer", "lid", "nr", "membernr", "memberno"},
    "name": {"naam", "name", "volledigenaam"},
    "phone": {"telefoon", "telnr", "telefoonnummer", "gsm", "gsmnr", "phone", "nummer", "tel"},
    "note": {"opmerking", "opmerkingen", "note", "notes"},
}

TRUE_WORDS = {"ja", "j", "yes", "y", "true", "waar", "1", "x", "aan", "on"}
FALSE_WORDS = {"nee", "neen", "n", "no", "false", "onwaar", "0", "uit", "off", "-"}


class ValidationError(ValueError):
    def __init__(self, errors):
        super().__init__("; ".join(errors))
        self.errors = errors


def _key(h: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def parse_bool(v):
    """True/False, of None voor leeg. Gooit ValueError bij onbekende waarde."""
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    s = str(v).strip().lower()
    if s.endswith(".0"):
        s = s[:-2]
    if s == "":
        return None
    if s in TRUE_WORDS:
        return True
    if s in FALSE_WORDS:
        return False
    raise ValueError(f"onbekende ja/nee-waarde: {v!r}")


def _clean_text(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s


def _access_map(member_ids: list[int] | None = None) -> dict[int, set[int]]:
    rows = db.query("SELECT member_id, device_id FROM member_access")
    out: dict[int, set[int]] = {}
    for r in rows:
        out.setdefault(r["member_id"], set()).add(r["device_id"])
    return out


def to_api(row: dict, access: set[int], devs: list[dict]) -> dict:
    out = {k: row[k] for k in ("id", *FIELDS, "created_at", "updated_at")}
    out["access"] = {d["id"]: d["id"] in access for d in devs}
    return out


def access_label(dev: dict) -> str:
    return f"Toegang {dev['name']}"


# ---------- CRUD ----------

def _parse_access(raw, devs: list[dict], base: dict[int, bool] | None, errors: list[str]) -> dict[int, bool]:
    """raw: {"<id>": bool}. Ontbrekende toestellen: waarde uit base, anders ja."""
    raw = raw or {}
    out = {}
    for d in devs:
        v = raw.get(str(d["id"]), raw.get(d["id"]))
        if v is None:
            out[d["id"]] = True if base is None else base.get(d["id"], True)
            continue
        try:
            b = parse_bool(v)
        except ValueError as e:
            errors.append(f"{access_label(d)}: {e}")
            b = None
        out[d["id"]] = True if b is None else b
    return out


def validate(data: dict, region: str, existing_id: int | None = None,
             base_access: dict[int, bool] | None = None) -> dict:
    errors = []
    out = {}
    out["member_no"] = _clean_text(data.get("member_no"))
    if not out["member_no"]:
        errors.append("Lidnr is verplicht.")
    out["name"] = _clean_text(data.get("name"))
    if not out["name"]:
        errors.append("Naam is verplicht.")
    try:
        out["phone"] = normalize(data.get("phone"), region)
    except PhoneError as e:
        errors.append(_cap(str(e)) + ".")
    out["note"] = _clean_text(data.get("note"))
    out["access"] = _parse_access(data.get("access"), devices.list_devices(), base_access, errors)
    if out["member_no"]:
        dup = db.query_one("SELECT id, member_no, name FROM members WHERE member_no=? AND id IS NOT ?",
                           (out["member_no"], existing_id))
        if dup:
            errors.append(f"Lidnr {out['member_no']} bestaat al ({dup['name']}).")
    if out.get("phone"):
        dup = db.query_one("SELECT member_no, name FROM members WHERE phone=? AND id IS NOT ?",
                           (out["phone"], existing_id))
        if dup:
            errors.append(f"Telefoon {out['phone']} hoort al bij lid {dup['member_no']} ({dup['name']}).")
    if errors:
        raise ValidationError(errors)
    return out


def list_members() -> list[dict]:
    devs = devices.list_devices()
    acc = _access_map()
    return [to_api(r, acc.get(r["id"], set()), devs)
            for r in db.query("SELECT * FROM members ORDER BY name COLLATE NOCASE, member_no")]


def get(member_id: int) -> dict | None:
    r = db.query_one("SELECT * FROM members WHERE id=?", (member_id,))
    if not r:
        return None
    acc = {x["device_id"] for x in db.query("SELECT device_id FROM member_access WHERE member_id=?", (member_id,))}
    return to_api(r, acc, devices.list_devices())


def find_by_phone(phone: str) -> dict | None:
    r = db.query_one("SELECT id FROM members WHERE phone=?", (phone,))
    return get(r["id"]) if r else None


def has_access(member: dict, device_id: int) -> bool:
    return bool(member["access"].get(device_id))


def _write_access(c, member_id: int, access: dict[int, bool]) -> None:
    for dev_id, allowed in access.items():
        if allowed:
            c.execute("INSERT OR IGNORE INTO member_access (member_id, device_id) VALUES (?,?)", (member_id, dev_id))
        else:
            c.execute("DELETE FROM member_access WHERE member_id=? AND device_id=?", (member_id, dev_id))


def create(data: dict, region: str, actor: str) -> dict:
    m = validate(data, region)
    ts = db.now_iso()
    with db.tx() as c:
        cur = c.execute("""INSERT INTO members (member_no, name, phone, note, created_at, updated_at)
                           VALUES (?,?,?,?,?,?)""", (m["member_no"], m["name"], m["phone"], m["note"], ts, ts))
        new_id = cur.lastrowid
        _write_access(c, new_id, m["access"])
    created = get(new_id)
    db.log_event("admin", actor=actor, member=created, phone=created["phone"], result="ok",
                 reason="Lid toegevoegd", detail=_describe(created))
    return created


def update(member_id: int, data: dict, region: str, actor: str) -> dict:
    old = get(member_id)
    if not old:
        raise KeyError(member_id)
    merged = {**old, **{k: v for k, v in data.items() if k in FIELDS}, "access": data.get("access")}
    m = validate(merged, region, existing_id=member_id, base_access=old["access"])
    with db.tx() as c:
        c.execute("UPDATE members SET member_no=?, name=?, phone=?, note=?, updated_at=? WHERE id=?",
                  (m["member_no"], m["name"], m["phone"], m["note"], db.now_iso(), member_id))
        _write_access(c, member_id, m["access"])
    new = get(member_id)
    changes = _changes(old, new)
    if changes:
        db.log_event("admin", actor=actor, member=new, phone=new["phone"], result="ok",
                     reason="Lid gewijzigd", detail="; ".join(changes))
    return new


def delete(member_id: int, actor: str) -> None:
    old = get(member_id)
    if not old:
        raise KeyError(member_id)
    db.execute("DELETE FROM members WHERE id=?", (member_id,))
    db.log_event("admin", actor=actor, member=old, phone=old["phone"], result="ok",
                 reason="Lid verwijderd", detail=_describe(old))


def _fmt(v):
    if isinstance(v, bool):
        return "ja" if v else "nee"
    return v if v not in (None, "") else "(leeg)"


def _pairs(m: dict, devs: list[dict]) -> list[tuple[str, object]]:
    out = [(LABELS[f], m[f]) for f in ("member_no", "name", "phone")]
    out += [(access_label(d), bool(m["access"].get(d["id"], False))) for d in devs]
    out.append((LABELS["note"], m["note"]))
    return out


def _describe(m: dict) -> str:
    return "; ".join(f"{k}: {_fmt(v)}" for k, v in _pairs(m, devices.list_devices()))


def _changes(old: dict, new: dict) -> list[str]:
    devs = devices.list_devices()
    return [f"{k}: {_fmt(a)} -> {_fmt(b)}"
            for (k, a), (_, b) in zip(_pairs(old, devs), _pairs(new, devs)) if a != b]


# ---------- export ----------

def export_rows() -> list[list]:
    devs = devices.list_devices()
    rows = [["Lidnr", "Naam", "Telefoon", *[access_label(d) for d in devs], "Opmerking"]]
    for m in list_members():
        rows.append([m["member_no"], m["name"], m["phone"] or "",
                     *[("ja" if m["access"].get(d["id"]) else "nee") for d in devs], m["note"] or ""])
    return rows


def export_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leden"
    rows = export_rows()
    for r in rows:
        ws.append(r)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = "@"   # tekst, zodat Excel +32... en lidnrs niet omvormt
    n = len(rows[0])
    widths = [10, 30, 16] + [max(14, len(h) + 2) for h in rows[0][3:n - 1]] + [40]
    for i, wdt in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = wdt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- import ----------

def _read_table(filename: str, content: bytes) -> list[list]:
    if not (filename or "").lower().endswith(".xlsx"):
        raise ValidationError(["Alleen Excel-bestanden (.xlsx) worden aanvaard, in het formaat van de export."])
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise ValidationError([f"Excel-bestand kan niet gelezen worden: {e}"]) from e
    ws = wb.worksheets[0]
    return [[("" if c is None else c) for c in row] for row in ws.iter_rows(values_only=True)]


def import_file(filename: str, content: bytes, region: str, actor: str) -> dict:
    """Upsert op lidnr. Leden die niet in het bestand staan, blijven behouden.
    Bij één fout wordt het hele bestand geweigerd."""
    devs = devices.list_devices()
    dev_keys = {}
    for d in devs:
        dev_keys["toegang" + _key(d["name"])] = d["id"]
        dev_keys.setdefault(_key(d["name"]), d["id"])
    table = _read_table(filename, content)
    while table and not any(str(c).strip() for c in table[0]):
        table.pop(0)
    if not table:
        raise ValidationError(["Het bestand is leeg."])
    header = table[0]
    colmap: dict = {}
    unknown = []
    for i, h in enumerate(header):
        k = _key(h)
        if not k:
            continue
        field = next((f for f, al in HEADER_ALIASES.items() if k in al), None)
        if field is None and k in dev_keys:
            field = ("access", dev_keys[k])
        if field is None:
            unknown.append(str(h))
        elif field in colmap:
            raise ValidationError([f"Kolom '{h}' komt dubbel voor."])
        else:
            colmap[field] = i
    if "member_no" not in colmap:
        raise ValidationError(["Kolom 'Lidnr' ontbreekt."])
    if unknown:
        allowed = ["Lidnr", "Naam", "Telefoon", *[access_label(d) for d in devs], "Opmerking"]
        raise ValidationError([f"Onbekende kolom(men): {', '.join(unknown)}. Toegelaten: {', '.join(allowed)}."])

    existing = {m["member_no"]: m for m in list_members()}
    errors: list[str] = []
    planned: dict[str, dict] = {}
    rownum_of: dict[str, int] = {}
    dev_name = {d["id"]: d["name"] for d in devs}

    for idx, row in enumerate(table[1:], start=2):
        cells = {f: (row[i] if i < len(row) else "") for f, i in colmap.items()}
        if not any(str(v).strip() for v in cells.values()):
            continue
        prefix = f"Rij {idx}: "
        member_no = _clean_text(cells["member_no"])
        if not member_no:
            errors.append(prefix + "Lidnr is leeg.")
            continue
        if member_no in rownum_of:
            errors.append(prefix + f"Lidnr {member_no} staat ook op rij {rownum_of[member_no]}.")
            continue
        rownum_of[member_no] = idx
        base = existing.get(member_no)
        target = ({**base, "access": dict(base["access"])} if base else
                  {"member_no": member_no, "name": "", "phone": None, "note": "",
                   "access": {d["id"]: True for d in devs}})
        if "name" in cells:
            name = _clean_text(cells["name"])
            if name:
                target["name"] = name
        if not target["name"]:
            errors.append(prefix + "Naam is verplicht voor een nieuw lid.")
        if "phone" in cells:
            try:
                target["phone"] = normalize(cells["phone"], region)
            except PhoneError as e:
                errors.append(prefix + _cap(str(e)) + ".")
        if "note" in cells:
            target["note"] = _clean_text(cells["note"])
        for f, v in cells.items():
            if isinstance(f, tuple):
                try:
                    b = parse_bool(v)
                except ValueError as e:
                    errors.append(prefix + f"Toegang {dev_name[f[1]]}: {e}.")
                    continue
                if b is not None:
                    target["access"][f[1]] = b
        planned[member_no] = target

    final = {**existing, **planned}
    seen: dict[str, str] = {}
    for no, m in final.items():
        p = m.get("phone")
        if not p:
            continue
        if p in seen:
            a, b = seen[p], no
            where = ", ".join(f"rij {rownum_of[x]}" if x in rownum_of else f"bestaand lid {x}" for x in (a, b))
            errors.append(f"Telefoon {p} komt twee keer voor: lid {a} en lid {b} ({where}).")
        else:
            seen[p] = no

    if errors:
        raise ValidationError(errors)

    created = changed = unchanged = 0
    per_member: list[tuple[str, dict, str]] = []
    ts = db.now_iso()
    with db.tx() as c:
        for no, t in planned.items():
            b = existing.get(no)
            if b and b["phone"] != t["phone"]:
                c.execute("UPDATE members SET phone=NULL WHERE id=?", (b["id"],))
        for no, t in planned.items():
            b = existing.get(no)
            if b is None:
                cur = c.execute("""INSERT INTO members (member_no, name, phone, note, created_at, updated_at)
                                   VALUES (?,?,?,?,?,?)""", (no, t["name"], t["phone"], t["note"], ts, ts))
                _write_access(c, cur.lastrowid, t["access"])
                created += 1
                per_member.append(("Import: lid toegevoegd", t, None))
            else:
                diff = _changes(b, t)
                if diff:
                    c.execute("UPDATE members SET name=?, phone=?, note=?, updated_at=? WHERE id=?",
                              (t["name"], t["phone"], t["note"], ts, b["id"]))
                    _write_access(c, b["id"], t["access"])
                    changed += 1
                    per_member.append(("Import: lid gewijzigd", t, "; ".join(diff)))
                else:
                    unchanged += 1

    for reason, t, detail in per_member:
        db.log_event("admin", actor=actor, member=t, phone=t["phone"], result="ok",
                     reason=reason, detail=detail or _describe(t))
    summary = {"nieuw": created, "gewijzigd": changed, "ongewijzigd": unchanged,
               "niet_in_bestand": len(set(existing) - set(planned))}
    db.log_event("admin", actor=actor, result="ok", reason="Import leden",
                 detail=f"{filename}: {created} nieuw, {changed} gewijzigd, {unchanged} ongewijzigd, "
                        f"{summary['niet_in_bestand']} niet in bestand (behouden)")
    return summary


def _cap(s: str) -> str:
    return s[:1].upper() + s[1:]
